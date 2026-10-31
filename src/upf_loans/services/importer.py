"""Import from the existing Excel workbook (MAIN calculation sheet layout).

The MAIN sheet has headers in row 3 and one employee per row:
A NAME, B EMPNO, C SALARY, D SAL_BAL, E PF balance, F contributions since,
H PF loan balance, I deductions since, K-P other loans, Q NTOT_LOBAL/SURETIES,
AI NIC, AJ/AK last loan month/year, AL DEPT, AO ACA OR NON ACA, AP DESIGNATION.
Values are read as Excel last calculated them, so cells such as
=37925+17800+17066 import as their result.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.eligibility import LoanInputs
from ..core.money import MoneyError, to_decimal
from ..db.models import Employee, User
from . import audit, employees, renewals
from .errors import ServiceError

COLUMNS = {
    "name": "A",
    "pf_number": "B",
    "salary": "C",
    "salary_balance": "D",
    "pf_balance": "E",
    "contributions_since": "F",
    "loan_balance": "H",
    "loan_deductions_since": "I",
    "distress_loan": "K",
    "staff_loan": "L",
    "vehicle_loan": "M",
    "computer_loan": "N",
    "other_loan_1": "O",
    "other_loan_2": "P",
    "q_amount": "Q",
    "nic": "AI",
    "last_loan_month": "AJ",
    "last_loan_year": "AK",
    "department": "AL",
    "category": "AO",
    "designation": "AP",
}
MONEY_FIELDS = (
    "salary", "salary_balance", "pf_balance", "contributions_since", "loan_balance", "loan_deductions_since",
    "distress_loan", "staff_loan", "vehicle_loan", "computer_loan", "other_loan_1", "other_loan_2", "q_amount",
)


@dataclass
class ImportRow:
    row: int
    values: dict
    errors: list[str] = field(default_factory=list)


@dataclass
class ImportReport:
    rows: list[ImportRow]
    created: int = 0
    updated: int = 0
    entries: int = 0

    @property
    def errors(self) -> list[str]:
        return [f"Row {r.row}: {e}" for r in self.rows for e in r.errors]


def _int_or_none(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(str(value).strip().split(".")[0])
    except ValueError:
        return None


def read_main_sheet(path: Path, sheet_name: str | None = None, first_row: int = 4) -> list[ImportRow]:
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name is None:
        candidates = [n for n in wb.sheetnames if "MAIN" in n.upper()]
        sheet_name = candidates[0] if candidates else wb.sheetnames[0]
    if sheet_name not in wb.sheetnames:
        raise ServiceError(f"Sheet '{sheet_name}' not found.")
    ws = wb[sheet_name]
    idx = {k: column_index_from_string(v) - 1 for k, v in COLUMNS.items()}
    rows: list[ImportRow] = []
    seen: dict[str, int] = {}
    for n, cells in enumerate(ws.iter_rows(min_row=first_row, values_only=True), start=first_row):
        cells = list(cells) + [None] * 50
        raw = {k: cells[i] for k, i in idx.items()}
        if not raw["pf_number"] or not str(raw["pf_number"]).strip():
            continue
        item = ImportRow(row=n, values={})
        for key, value in raw.items():
            if key in MONEY_FIELDS:
                try:
                    item.values[key] = to_decimal(value)
                except MoneyError:
                    item.errors.append(f"{key} '{value}' is not a number")
            elif key in ("last_loan_month", "last_loan_year"):
                item.values[key] = _int_or_none(value)
            else:
                item.values[key] = "" if value is None else str(value).strip()
        if not item.values.get("category"):
            item.values["category"] = "N"
        data = employees.EmployeeData(
            pf_number=item.values["pf_number"], name=item.values["name"], nic=item.values["nic"],
            department=item.values["department"], category=item.values["category"],
            designation=item.values["designation"],
        ).cleaned()
        if data.pf_number in seen:
            item.errors.append(f"PF number {data.pf_number} is repeated (also row {seen[data.pf_number]})")
        seen.setdefault(data.pf_number, n)
        item.values["employee"] = data
        rows.append(item)
    wb.close()
    return rows


def import_rows(session: Session, actor: User, rows: list[ImportRow], batch_id: int | None) -> ImportReport:
    """Import valid rows; rows with errors are skipped and reported."""
    report = ImportReport(rows=rows)
    for item in rows:
        if item.errors:
            continue
        data: employees.EmployeeData = item.values["employee"]
        existing = session.scalar(select(Employee).where(Employee.pf_number == data.pf_number))
        errs = employees.validate(session, data, existing.id if existing else None)
        if errs:
            item.errors.extend(errs.values())
            continue
        emp = employees.save(session, actor, data, existing.id if existing else None)
        if existing:
            report.updated += 1
        else:
            report.created += 1
        if batch_id is not None:
            inputs = LoanInputs(**{k: item.values[k] for k in MONEY_FIELDS})
            try:
                renewals.save_entry(
                    session, actor, batch_id, emp.id, inputs,
                    item.values["last_loan_month"], item.values["last_loan_year"],
                )
                report.entries += 1
            except ServiceError as exc:
                item.errors.append(str(exc))
    audit.record(
        session, actor.username, "import.excel", "employee",
        summary=f"{report.created} created, {report.updated} updated, {report.entries} calculations, {len(report.errors)} errors",
    )
    return report
