"""UGC renewal sheet, reproducing the "PF LOAN send UGC" sheet layout."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins

from ..core.ugc_text import UgcRow
from ..services.renewals import month_name

HEADERS = [
    "", "NAME", "PF-NUMBER", "ID-NUMBER", "LOAN REQUEST MONTH", "LOAN REQUEST YEAR", "LAST LOAN MONTH",
    "LAST LOAN YEAR", "LOAN APPROVED", "EXISTING LOAN BALANCE", "BALANCE  PAYABLE", "UNIVERSITY CODE",
]
WIDTHS = {"A": 3.4, "B": 36.6, "C": 6.7, "D": 10.6, "E": 2.1, "F": 3.7, "G": 2.2, "H": 3.9,
          "I": 12.9, "J": 12.4, "K": 12.3, "L": 2.4}
AMOUNT_FORMAT = "000000000000.00"
THIN = Side(style="thin")
MEDIUM = Side(style="medium")
COLS = "ABCDEFGHIJKL"


def title_text(month: int) -> str:
    return f"RENEWAL OF U.P.F. LOAN - APPLICATION FOR THE MONTH OF  {month_name(month)}"


def write_ugc_workbook(
    path: Path, rows: list[UgcRow], request_month: int, university_name: str, signatory: str, letter_date: str
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "PF LOAN send UGC"
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width

    title_font = Font(name="Bookman Old Style", size=12, bold=True)
    head_font = Font(name="Bookman Old Style", size=9, bold=True)
    body_font = Font(name="Arial", size=9)
    bold_font = Font(name="Arial", size=9, bold=True)
    center = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("B1:J1")
    ws.merge_cells("B2:J2")
    ws["B1"] = title_text(request_month)
    ws["B2"] = university_name
    for ref in ("B1", "B2"):
        ws[ref].font = title_font
        ws[ref].alignment = center
    ws.row_dimensions[1].height = 14.25
    ws.row_dimensions[2].height = 12.75

    ws.row_dimensions[3].height = 35.25
    for i, text in enumerate(HEADERS):
        cell = ws.cell(row=3, column=i + 1, value=text or None)
        cell.font = head_font
        left = COLS[i] in "EFGH"
        cell.alignment = Alignment(horizontal="left" if left else "center", vertical="center", wrap_text=True)

    r = 4
    for n, row in enumerate(rows, start=1):
        values = [
            n, row.name, row.pf_number, row.nic, f"{row.request_month:02d}", row.request_year,
            f"{row.last_loan_month:02d}" if row.last_loan_month else None,
            str(row.last_loan_year) if row.last_loan_year else None,
            float(row.loan_approved), float(row.existing_balance), float(row.balance_payable), row.university_code,
        ]
        for i, value in enumerate(values):
            cell = ws.cell(row=r, column=i + 1, value=value)
            cell.font = body_font
            if COLS[i] in "IJK":
                cell.number_format = AMOUNT_FORMAT
            elif COLS[i] in "EGHL":
                cell.number_format = "@"
        ws.row_dimensions[r].height = 12
        r += 1

    total_row = r
    for col in "IJK":
        cell = ws[f"{col}{total_row}"]
        cell.value = f"=SUM({col}4:{col}{total_row - 1})" if rows else 0
        cell.font = bold_font
        cell.number_format = AMOUNT_FORMAT
        cell.border = Border(top=THIN, bottom=THIN)

    sig = total_row + 3
    for offset, text in enumerate((f" {signatory} ", f" {university_name} ", letter_date)):
        ws[f"B{sig + offset}"] = text
        ws[f"B{sig + offset}"].font = body_font
    last = sig + 2

    # Box border round the whole form, as in the original sheet.
    for row_idx in range(1, last + 1):
        for col in COLS:
            cell = ws[f"{col}{row_idx}"]
            b = cell.border
            cell.border = Border(
                left=MEDIUM if col == "A" else b.left,
                right=MEDIUM if col == "L" else b.right,
                top=MEDIUM if row_idx == 1 else b.top,
                bottom=MEDIUM if row_idx == last else b.bottom,
            )

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "3:3"
    ws.print_area = f"A1:L{last}"
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    path = Path(path)
    wb.save(path)
    return path
