"""Calculation sheet export in the MAIN sheet layout, with live Excel formulas."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..core.eligibility import LoanRules
from ..db.models import RenewalBatch

MONEY = '_-* #,##0.00_-;\\-* #,##0.00_-;_-* "-"??_-;_-@_-'
INPUT_HEADERS = {
    "A": "NAME", "B": "EMPNO", "C": "SALARY", "D": "SAL_BAL", "E": "PF BAL", "F": "CON SINCE",
    "G": "PF BALANCE", "H": "PFLBAL", "I": "PFLDED", "J": "PFLBAL NOW", "K": "AU_DIS", "L": "STAFF LOAN",
    "M": "AU_VEH", "N": "COMPUTER", "O": "Loan1", "P": "Loan2", "Q": "NTOT_LOBAL/SURETIES",
}


def write_calculation_workbook(path: Path, batch: RenewalBatch, rules: LoanRules, title: str) -> Path:
    sm, bm, pr, sr, n = (rules.salary_multiplier, rules.salary_balance_multiplier, rules.pf_balance_ratio,
                         rules.surety_allowance_ratio, rules.instalment_months)
    result_headers = {
        "S": "NAME", "T": "EMPNO", "U": f"SAL*{sm}", "V": f"SAL_BAL*{bm}", "W": f"BAL*{pr}", "X": "LOWEST",
        "Y": "LOAN ABILITY", "Z": "INS", "AA": "ELIGIBLE", "AB": "LOAN_BAL", "AC": "PAY LOAN",
        "AD": "IF INS=SALBAL", "AE": "EMPNO", "AF": "INS.", "AG": "LOAN IN HAND",
        "AI": "ID. NO", "AJ": "LAST LOAN MONTH", "AK": "LAST LOAN YEAR", "AL": "DEPT",
        "AO": "ACA OR NON ACA", "AP": "DESIGNATION",
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "MAIN"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12)
    input_fill = PatternFill("solid", fgColor="FFF7E0")
    for col, text in {**INPUT_HEADERS, **result_headers}.items():
        c = ws[f"{col}3"]
        c.value = text
        c.font = Font(bold=True, size=9)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A4"
    for i, e in enumerate(batch.entries):
        r = 4 + i
        emp = e.employee
        values = {
            "A": emp.name, "B": emp.pf_number, "C": e.salary, "D": e.salary_balance, "E": e.pf_balance,
            "F": e.contributions_since, "H": e.loan_balance, "I": e.loan_deductions_since, "K": e.distress_loan,
            "L": e.staff_loan, "M": e.vehicle_loan, "N": e.computer_loan, "O": e.other_loan_1,
            "P": e.other_loan_2, "Q": e.q_amount,
        }
        for col, v in values.items():
            cell = ws[f"{col}{r}"]
            cell.value = float(v) if not isinstance(v, str) else v
            cell.fill = input_fill
        formulas = {
            "G": f"=E{r}+F{r}", "J": f"=H{r}-I{r}", "S": f"=A{r}", "T": f"=B{r}",
            "U": f"=C{r}*{sm}", "V": f"=D{r}*{bm}", "W": f"=G{r}*{pr}", "X": f"=MIN(U{r}:W{r})",
            "Y": f"=IF(Q{r}>(G{r}*{sr}),(X{r}-(Q{r}-(G{r}*{sr}))),X{r})", "Z": f"=INT(Y{r}/{n})",
            "AA": f"=Z{r}*{n}", "AB": f"=J{r}", "AC": f"=AA{r}-AB{r}",
            "AD": f'=IF(D{r}<=Z{r},"< 60%"," ")', "AE": f"=T{r}", "AF": f"=AA{r}/{n}",
            "AG": f"=AC{r}-(K{r}+L{r}+M{r}+N{r}+O{r}+P{r})",
        }
        for col, f in formulas.items():
            ws[f"{col}{r}"] = f
        extra = {
            "AI": emp.nic, "AJ": f"{e.last_loan_month:02d}" if e.last_loan_month else None,
            "AK": str(e.last_loan_year) if e.last_loan_year else None, "AL": emp.department,
            "AO": emp.category, "AP": emp.designation,
        }
        for col, v in extra.items():
            ws[f"{col}{r}"] = v
    for col in list(INPUT_HEADERS) + list(result_headers):
        ws.column_dimensions[col].width = 14 if col not in ("A", "S", "AP") else 24
        if col not in ("A", "B", "S", "T", "AD", "AE", "AI", "AJ", "AK", "AL", "AO", "AP"):
            for r in range(4, 4 + len(batch.entries)):
                ws[f"{col}{r}"].number_format = MONEY
    path = Path(path)
    wb.save(path)
    return path
