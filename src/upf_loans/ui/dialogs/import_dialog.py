"""Import employees (and optionally their calculation inputs) from the Excel MAIN sheet."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtWidgets import (
    QComboBox, QDialog, QFileDialog, QFormLayout, QHBoxLayout, QLineEdit, QVBoxLayout, QWidget,
)

from ...services import importer, renewals
from ...services.errors import ServiceError
from ..context import AppContext
from ..widgets.common import Banner, button, error, info, muted
from ..widgets.table import Column, DataTable
from .. import theme


class ImportDialog(QDialog):
    def __init__(self, ctx: AppContext, batch_id: int | None = None, parent: QWidget | None = None):
        super().__init__(parent)
        self.ctx = ctx
        self.rows: list[importer.ImportRow] = []
        self.setWindowTitle("Import from Excel")
        self.resize(980, 600)
        lay = QVBoxLayout(self)
        lay.addWidget(muted(
            "Pick a workbook laid out like the MAIN calculation sheet (headers in row 3, one employee per row). "
            "Employees are matched by PF number: new ones are added, existing ones are updated."
        ))
        form = QFormLayout()
        file_row = QHBoxLayout()
        self.path = QLineEdit()
        self.path.setReadOnly(True)
        browse = button("Browse…", "file")
        browse.clicked.connect(self._browse)
        file_row.addWidget(self.path, 1)
        file_row.addWidget(browse)
        form.addRow("Workbook", file_row)
        self.sheet = QComboBox()
        self.sheet.currentTextChanged.connect(lambda _t: self._preview())
        form.addRow("Sheet", self.sheet)
        self.batch = QComboBox()
        self.batch.addItem("Employees only (no calculations)", None)
        with ctx.db.session() as s:
            for b in renewals.list_batches(s):
                if not b.is_locked:
                    self.batch.addItem(f"Also add calculations to {renewals.batch_label(b)}", b.id)
        if batch_id is not None:
            idx = self.batch.findData(batch_id)
            if idx >= 0:
                self.batch.setCurrentIndex(idx)
        form.addRow("Calculations", self.batch)
        lay.addLayout(form)
        self.banner = Banner()
        lay.addWidget(self.banner)
        self.table = DataTable([
            Column("Row", "row", "int", width=55), Column("PF number", "pf", width=100), Column("Name", "name", stretch=True),
            Column("NIC", "nic", width=120), Column("Salary", "salary", "money", width=110),
            Column("PF balance", "pf_bal", "money", width=120), Column("Problem", "problem", width=260,
                                                                       color=lambda r: theme.DANGER if r["problem"] else None),
        ])
        lay.addWidget(self.table, 1)
        buttons = QHBoxLayout()
        buttons.addStretch(1)
        cancel = button("Cancel")
        cancel.clicked.connect(self.reject)
        self.go = button("Import", "upload", "primary")
        self.go.setEnabled(False)
        self.go.clicked.connect(self._import)
        buttons.addWidget(cancel)
        buttons.addWidget(self.go)
        lay.addLayout(buttons)

    def _browse(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "Choose workbook", str(Path.home()), "Excel workbook (*.xlsx *.xlsm)")
        if not path:
            return
        self.path.setText(path)
        try:
            wb = load_workbook(path, read_only=True)
            names = wb.sheetnames
            wb.close()
        except Exception as exc:  # noqa: BLE001 - any unreadable file
            error(self, f"Could not open the workbook:\n{exc}")
            return
        self.sheet.blockSignals(True)
        self.sheet.clear()
        self.sheet.addItems(names)
        main = [n for n in names if "MAIN" in n.upper()]
        if main:
            self.sheet.setCurrentText(main[0])
        self.sheet.blockSignals(False)
        self._preview()

    def _preview(self) -> None:
        if not self.path.text() or not self.sheet.currentText():
            return
        try:
            self.rows = importer.read_main_sheet(Path(self.path.text()), self.sheet.currentText())
        except (ServiceError, OSError, ValueError) as exc:
            self.banner.show_message("error", str(exc))
            self.rows = []
        self.table.set_rows([
            {"row": r.row, "pf": r.values["employee"].pf_number, "name": r.values["employee"].name,
             "nic": r.values["employee"].nic, "salary": r.values.get("salary"), "pf_bal": r.values.get("pf_balance"),
             "problem": "; ".join(r.errors)}
            for r in self.rows
        ])
        bad = sum(1 for r in self.rows if r.errors)
        if not self.rows:
            self.banner.show_message("warning", "No employee rows found (column B, from row 4).")
        elif bad:
            self.banner.show_message("warning", f"{len(self.rows)} rows found; {bad} have problems and will be skipped.")
        else:
            self.banner.show_message("ok", f"{len(self.rows)} rows ready to import.")
        self.go.setEnabled(any(not r.errors for r in self.rows))

    def _import(self) -> None:
        try:
            with self.ctx.db.session() as s:
                report = importer.import_rows(s, self.ctx.operator, self.rows, self.batch.currentData())
        except ServiceError as exc:
            error(self, str(exc))
            return
        self.ctx.events.data_changed.emit("employees")
        msg = f"{report.created} employees added, {report.updated} updated, {report.entries} calculations saved."
        if report.errors:
            msg += "\n\nSkipped:\n" + "\n".join(report.errors[:15])
        info(self, msg, "Import finished")
        self.accept()
