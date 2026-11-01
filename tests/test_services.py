import re
from decimal import Decimal as D

import pytest
from openpyxl import Workbook, load_workbook

from upf_loans.core.eligibility import LoanInputs
from upf_loans.services import backup, employees, exports, importer, renewals, settings
from upf_loans.services.errors import ServiceError


def add_employee(s, admin, pf="R 001365", name="Perera A.B. Mr.", nic="680410054V"):
    return employees.save(s, admin, employees.EmployeeData(pf_number=pf, name=name, nic=nic))


ROW7 = LoanInputs(salary=D(72791), salary_balance=D(26455), pf_balance=D("3622669.53"),
                  contributions_since=D("105650.79"), loan_balance=D(1760970), loan_deductions_since=D(150500),
                  vehicle_loan=D(30000))


def test_employee_validation_and_duplicates(db, admin):
    with db.session() as s:
        e = add_employee(s, admin, pf="r001365")
        assert e.pf_number == "R001365" and e.name == "PERERA A.B. MR."
        with pytest.raises(ServiceError, match="already belongs"):
            add_employee(s, admin, pf="R001365")
        with pytest.raises(ServiceError, match="NIC"):
            add_employee(s, admin, pf="R 000002", nic="12345")
        with pytest.raises(ServiceError, match="10 characters"):
            add_employee(s, admin, pf="R 000003", nic="197270000375")
        assert employees.search(s, "perera")[0].id == e.id
        assert employees.search(s, "680410054")[0].id == e.id


def test_batch_entry_money_round_trip_and_lock(db, admin):
    with db.session() as s:
        emp = add_employee(s, admin)
        batch = renewals.create_batch(s, admin, 11, 2024)
        entry, result = renewals.save_entry(s, admin, batch.id, emp.id, ROW7, 10, 2023)
        assert result.balance_payable == D(812930)
        batch_id = batch.id
        with pytest.raises(ServiceError, match="already exists"):
            renewals.create_batch(s, admin, 11, 2024)
    with db.session() as s:
        b = renewals.get_batch(s, batch_id)
        e = b.entries[0]
        assert e.contributions_since == D("105650.79")
        assert e.loan_approved == D(2423400) and e.loan_in_hand == D(782930)
        renewals.finalize(s, admin, batch_id, "14.11.2024")
    with db.session() as s:
        with pytest.raises(ServiceError, match="finalized"):
            renewals.save_entry(s, admin, batch_id, e.employee_id, ROW7, 10, 2023)


def test_saving_same_employee_twice_updates_not_duplicates(db, admin):
    with db.session() as s:
        emp = add_employee(s, admin)
        batch = renewals.create_batch(s, admin, 11, 2024)
        renewals.save_entry(s, admin, batch.id, emp.id, ROW7, 10, 2023)
        renewals.save_entry(s, admin, batch.id, emp.id, ROW7, 9, 2023)
        assert len(batch.entries) == 1 and batch.entries[0].last_loan_month == 9


def test_last_loan_suggested_from_finalized_month(db, admin):
    with db.session() as s:
        emp = add_employee(s, admin)
        b1 = renewals.create_batch(s, admin, 11, 2024)
        renewals.save_entry(s, admin, b1.id, emp.id, ROW7, 10, 2023)
        renewals.finalize(s, admin, b1.id, "14.11.2024")
        b2 = renewals.create_batch(s, admin, 12, 2025)
        assert renewals.suggested_last_loan(s, emp.id, b2) == (11, 2024)


def test_reopen_requires_reason(db, admin):
    with db.session() as s:
        emp = add_employee(s, admin)
        b = renewals.create_batch(s, admin, 11, 2024)
        renewals.save_entry(s, admin, b.id, emp.id, ROW7, 10, 2023)
        renewals.finalize(s, admin, b.id, "14.11.2024")
        with pytest.raises(ServiceError):
            renewals.reopen(s, admin, b.id, " ")
        assert renewals.reopen(s, admin, b.id, "typo").status == "draft"


def test_settings_change_rules_used_for_new_calculations(db, admin):
    with db.session() as s:
        settings.save(s, admin, {"salary_multiplier": "24"})
        assert settings.loan_rules(s).salary_multiplier == D(24)
        with pytest.raises(ServiceError):
            settings.save(s, admin, {"pf_balance_ratio": "1.5"})
        with pytest.raises(ServiceError):
            settings.save(s, admin, {"salary_multiplier": "abc"})


def test_exports(db, admin, tmp_path):
    with db.session() as s:
        e1 = add_employee(s, admin)
        e2 = add_employee(s, admin, pf="F000288", name="Silva C.", nic="677910690V")
        b = renewals.create_batch(s, admin, 11, 2024)
        renewals.save_entry(s, admin, b.id, e1.id, ROW7, 10, 2023)
        renewals.save_entry(s, admin, b.id, e2.id, ROW7, 7, 2023)
        txt = exports.export_ugc_text(s, admin, b.id, tmp_path / "out.txt")
        lines = txt.read_bytes().decode("cp1252").split("\r\n")
        # Same as Excel's "Text (Tab delimited)" save of the PF LOAN send UGC sheet
        assert lines[0] == "\tRENEWAL OF U.P.F. LOAN - APPLICATION FOR THE MONTH OF  NOVEMBER"
        assert lines[1] == "\tUNIVERSITY OF RUHUNA"
        assert lines[2].split("\t") == ["", "NAME", "PF-NUMBER", "ID-NUMBER", "LOAN REQUEST MONTH",
                                        "LOAN REQUEST YEAR", "LAST LOAN MONTH", "LAST LOAN YEAR", "LOAN APPROVED",
                                        "EXISTING LOAN BALANCE", "BALANCE  PAYABLE", "UNIVERSITY CODE"]
        assert lines[3].split("\t")[0] == "1"
        assert lines[3].split("\t")[2:] == ["R 001365", "680410054V", "11", "2024", "10", "2023",
                                            "000002423400.00", "000001610470.00", "000000812930.00", "F"]
        assert lines[4].split("\t")[2] == "F000288" and lines[4].split("\t")[6] == "07"
        assert lines[5] == "\t" * 8 + "000004846800.00\t000003220940.00\t000001625860.00"
        assert lines[6] == lines[7] == ""
        assert lines[8:10] == ["\t BURSAR ", "\t UNIVERSITY OF RUHUNA "]
        assert re.fullmatch(r"\t\d\d\.\d\d\.\d{4}", lines[10])  # letter date
        assert lines[11] == "" and len(lines) == 12
        xlsx = exports.export_ugc_excel(s, admin, b.id, tmp_path / "ugc.xlsx")
        ws = load_workbook(xlsx).active
        assert ws["B1"].value.endswith("NOVEMBER")
        assert ws["C4"].value == "R 001365" and ws["I4"].number_format == "000000000000.00"
        assert ws["I6"].value == "=SUM(I4:I5)"
        assert ws["B9"].value.strip() == "BURSAR"
        exports.export_calculation_excel(s, admin, b.id, tmp_path / "calc.xlsx")
        exports.export_pdf(s, admin, b.id, tmp_path / "r.pdf")
        assert (tmp_path / "r.pdf").read_bytes()[:4] == b"%PDF"


def test_text_export_blocks_on_bad_nic(db, admin, tmp_path):
    with db.session() as s:
        emp = add_employee(s, admin)
        emp.nic = "197270000375"  # e.g. data from before the 10-character rule
        b = renewals.create_batch(s, admin, 11, 2024)
        renewals.save_entry(s, admin, b.id, emp.id, ROW7, 10, 2023)
        with pytest.raises(ServiceError, match="NIC"):
            exports.export_ugc_text(s, admin, b.id, tmp_path / "x.txt")
        assert not (tmp_path / "x.txt").exists()


def make_main_sheet(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2024  MAIN  NOVEMBER"
    ws["A3"], ws["B3"] = "NAME", "EMPNO"
    rows = [
        (5, ["A", "R000001", 1000, 2000, 80000, 0, None, 500, 200, None, 0, 0, 0, 0, 0, 0, 12000], "680410054V"),
        (7, ["B", "R 001365", 72791, 26455, 3622669.53, 105650.79, None, 1760970, 150500, None, 0, 0, 30000, 0, 0, 0, 0], "680410055V"),
        (8, ["C", "R 001365", 1, 1, 1, 0, None, 0, 0, None, 0, 0, 0, 0, 0, 0, 0], "680410056V"),
    ]
    for r, values, nic in rows:
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
        ws[f"AI{r}"], ws[f"AJ{r}"], ws[f"AK{r}"], ws[f"AO{r}"] = nic, "10", "2023", "N"
    wb.save(path)


def test_import_main_sheet(db, admin, tmp_path):
    path = tmp_path / "main.xlsx"
    make_main_sheet(path)
    rows = importer.read_main_sheet(path)
    assert len(rows) == 3
    assert rows[2].errors and "repeated" in rows[2].errors[0]
    with db.session() as s:
        b = renewals.create_batch(s, admin, 11, 2024)
        report = importer.import_rows(s, admin, rows, b.id)
        assert report.created == 2 and report.entries == 2
        payable = {e.employee.pf_number: e.balance_payable for e in b.entries}
        assert payable == {"R000001": D(43620), "R 001365": D(812930)}


def test_backup_and_restore(db, admin, tmp_path):
    with db.session() as s:
        add_employee(s, admin)
    bdir = tmp_path / "backups"
    path = backup.create_backup(db.path, bdir)
    backup.validate_backup(path)
    with db.session() as s:
        add_employee(s, admin, pf="R 000002", nic="680410057V")
    db.dispose()
    backup.restore(path, db.path, bdir)
    from upf_loans.db.session import Database
    db2 = Database(db.path)
    with db2.session() as s:
        assert len(employees.search(s)) == 1
    db2.dispose()
    bad = tmp_path / "bad.sqlite3"
    bad.write_text("nope")
    with pytest.raises(ServiceError):
        backup.validate_backup(bad)


def test_move_data_folder(tmp_path, monkeypatch, admin):
    from upf_loans import config
    from upf_loans.db.session import Database
    from upf_loans.services import location

    monkeypatch.setenv("UPF_LOANS_HOME", str(tmp_path / "home"))
    db = Database(config.database_path())
    with db.session() as s:
        add_employee(s, admin)
    backup.create_backup(config.database_path(), config.backup_dir())
    db.dispose()

    target = tmp_path / "D" / "UPFLoans"
    location.move_data(target)
    assert config.data_dir() == target.resolve()
    assert config.database_path() == target.resolve() / "upf_loans.sqlite3"
    assert len(backup.list_backups(config.backup_dir())) == 1
    db2 = Database(config.database_path())
    with db2.session() as s:
        assert len(employees.search(s)) == 1
    db2.dispose()
    with pytest.raises(ServiceError, match="already"):
        location.move_data(target)

    # a missing drive gives a clear error instead of silently starting empty
    config.set_data_dir(tmp_path / "file.txt" / "x")
    (tmp_path / "file.txt").write_text("not a folder")
    with pytest.raises(config.DataLocationError):
        config.data_dir()
    config.set_data_dir(None)
    assert config.data_dir() == config.home_dir()
